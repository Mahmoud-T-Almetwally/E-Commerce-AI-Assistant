"""
Extracts plain text from uploaded knowledge-base files (.txt, .pdf, .docx).

Parsing happens from in-memory bytes — uploaded files are never written to disk.
"""

from io import BytesIO
from utils.exceptions import TextExtractionError


def extract_text(data: bytes, extension: str) -> str:
    """
    Extracts text from raw file bytes. `extension` must include the leading
    dot, e.g. ".pdf". Raises TextExtractionError on unsupported types,
    unreadable content, or missing parser libraries.
    """
    dispatch = {
        ".txt": _extract_txt,
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
    }
    handler = dispatch.get((extension or "").lower())
    if handler is None:
        raise TextExtractionError(
            f"Unsupported file type '{extension}'. Allowed: .txt, .pdf, .docx"
        )

    text = handler(data)
    if not text or not text.strip():
        raise TextExtractionError(
            "No text could be extracted from this file. It may be empty, "
            "image-only (e.g. a scanned PDF), or corrupted."
        )
    return text.strip()


def _extract_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise TextExtractionError(
            "PDF support is unavailable: the 'pypdf' package is not installed "
            "(pip install pypdf)."
        )
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise TextExtractionError("Password-protected PDFs are not supported.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n\n".join(pages)
    except TextExtractionError:
        raise
    except Exception as e:
        raise TextExtractionError(f"Could not read this PDF: {e}") from e


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError:
        raise TextExtractionError(
            "Word support is unavailable: the 'python-docx' package is not "
            "installed (pip install python-docx)."
        )
    try:
        doc = DocxDocument(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as e:
        raise TextExtractionError(f"Could not read this .docx file: {e}") from e